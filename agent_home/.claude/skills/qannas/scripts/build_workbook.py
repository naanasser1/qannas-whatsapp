#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""يبني ملف Excel لدفعة قناص من leads.json — أربعة تبويبات جاهزة لـ Google Sheets.

    python build_workbook.py leads.json --out qannas-2026-08-30.xlsx

التبويبات: E-commerce Leads · Upwork · Mostaql · المنهجية والمصادر.
شكل leads.json موصوف في scripts/README.md
"""
import argparse
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import qannas_lib as lib  # noqa: E402

INK = "1B1C1E"
ACCENT = "1F4E46"
HEAD_FILL = PatternFill("solid", fgColor=ACCENT)
ALT_FILL = PatternFill("solid", fgColor="F2F1ED")
MSG_FILL = PatternFill("solid", fgColor="FFF8E7")
HOT_FILL = PatternFill("solid", fgColor="FBE3DC")
HIGH_FILL = PatternFill("solid", fgColor="E2EFE6")
_thin = Side(style="thin", color="D6D2C8")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

TIER_FILL = {"🔥 HOT LEAD": HOT_FILL, "🟢 HIGH PRIORITY": HIGH_FILL}

# (العنوان، العرض) — الأعمدة الـ٢١ من المواصفة + عمود متابعة الحالة
LEAD_COLUMNS = [
    ("Date · التاريخ", 12),
    ("Priority · الأولوية", 16),
    ("Company · الشركة", 22),
    ("Country · الدولة", 12),
    ("Industry · المجال", 20),
    ("Website · الموقع", 26),
    ("Website Status · حالة الموقع", 24),
    ("UX/UI Issues · ملاحظات التجربة والواجهة", 46),
    ("Logo/Brand Issues · ملاحظات الهوية", 38),
    ("Image Issues · ملاحظات الصور", 38),
    ("Website Freshness · حداثة الموقع", 30),
    ("Ad/Video Issues · ملاحظات الإعلانات", 46),
    ("Primary Opportunity · الفرصة الأساسية", 30),
    ("Recommended Service · الخدمة المقترحة", 30),
    ("Score · الدرجة", 10),
    ("Decision Maker · صانع القرار", 26),
    ("Email", 26),
    ("LinkedIn", 28),
    ("Instagram", 20),
    ("Personalized Message · الرسالة الجاهزة", 62),
    ("Source · المصدر", 24),
    ("Status · الحالة", 16),
]

PROJECT_COLUMNS = [
    ("Date · التاريخ", 12),
    ("Priority · الأولوية", 14),
    ("Project · المشروع", 44),
    ("Client · العميل", 20),
    ("URL · الرابط", 44),
    ("Budget · الميزانية", 18),
    ("Posted Date · تاريخ النشر", 18),
    ("Skills · المهارات", 30),
    ("Summary · الملخص", 44),
    ("Why It Fits · لماذا يناسب", 40),
    ("Proposal Angle · زاوية العرض", 44),
    ("Score · الدرجة", 10),
    ("Status · الحالة", 16),
]

LEAD_STATUS = '"جاهزة للإرسال,تم الإرسال,رد إيجابي,لا يوجد رد,مغلق"'
PROJECT_STATUS = '"لم يُقدَّم بعد,تم التقديم,مغلق,غير مناسب"'


def _style_header(ws, count, height=38):
    for c in range(1, count + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = height


def _auto_height(texts, widths, min_height, line_pt=13.5, max_height=409):
    """ارتفاع صف يتسع فعلاً للنص.

    Excel لا يحسب الارتفاع تلقائياً مع wrap_text عند ضبطه من openpyxl، فالنص
    الطويل يُقتطع بصرياً. نقدّر عدد الأسطر من عدد الحروف مقسوماً على عرض العمود،
    ونأخذ أطول خلية في الصف. هذا مهم تحديداً لتبويب المنهجية: نص حدود التحقق هو
    أهم ما في الدفعة، وقصّه يعني أن ناصر لن يقرأ ما لم يُتحقق منه.
    """
    lines = 1
    for text, width in zip(texts, widths):
        chars_per_line = max(int(width * 1.05), 10)
        needed = sum(max(1, -(-len(p) // chars_per_line))
                     for p in str(text or "").split("\n"))
        lines = max(lines, needed)
    return min(max(min_height, lines * line_pt + 8), max_height)


def _style_body(ws, rows, cols, min_height, widths, skip_fill=()):
    for r in range(2, 2 + rows):
        texts = [ws.cell(row=r, column=c).value for c in range(1, cols + 1)]
        ws.row_dimensions[r].height = _auto_height(texts, widths, min_height)
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name="Arial", size=10, color=INK)
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            cell.border = BORDER
            if r % 2 == 0 and c not in skip_fill:
                cell.fill = ALT_FILL


def _apply_widths(ws, columns):
    for i, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _link(cell, size=9):
    if cell.value and str(cell.value).startswith("http"):
        cell.hyperlink = str(cell.value)
        cell.font = Font(name="Arial", size=size, color=ACCENT, underline="single")


def _dropdown(ws, column_letter, rows, formula):
    if not rows:
        return
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{column_letter}2:{column_letter}{1 + rows}")


def _leads_sheet(wb, leads, date):
    ws = wb.active
    ws.title = "E-commerce Leads"
    ws.sheet_view.rightToLeft = True
    ws.append([title for title, _ in LEAD_COLUMNS])
    _style_header(ws, len(LEAD_COLUMNS))

    for lead in leads:
        dm = lead["decision_maker"]
        ws.append([
            date, lead["tier"], lead["company"], lead["country"], lead["industry"],
            lead["website"], lead["website_status"],
            lead["ux_ui_text"], lead["brand_text"], lead["images_text"],
            lead["freshness"], lead["ads_text"],
            lead["primary_opportunity"], lead["recommended_service"],
            lead["total"], lead["dm_text"],
            dm["email"] or lib.UNVERIFIED,
            dm["linkedin"] or lib.UNVERIFIED,
            dm["instagram"] or lib.UNVERIFIED,
            lead["message"], lead["source"], "جاهزة للإرسال",
        ])

    widths = [w for _, w in LEAD_COLUMNS]
    _style_body(ws, len(leads), len(LEAD_COLUMNS), 120, widths, skip_fill=(20,))

    for i, lead in enumerate(leads):
        r = 2 + i
        ws.cell(row=r, column=2).fill = TIER_FILL.get(lead["tier"], ALT_FILL)
        ws.cell(row=r, column=2).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True)
        ws.cell(row=r, column=3).font = Font(name="Arial", size=10, bold=True, color=INK)
        score = ws.cell(row=r, column=15)
        score.font = Font(name="Arial", size=12, bold=True, color=ACCENT)
        score.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=20).fill = MSG_FILL
        for col in (6, 18, 19):
            _link(ws.cell(row=r, column=col))
        ws.cell(row=r, column=22).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True)

    _apply_widths(ws, LEAD_COLUMNS)
    _dropdown(ws, "V", len(leads), LEAD_STATUS)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:V{1 + len(leads)}"
    return ws


def _projects_sheet(wb, title, projects, date):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append([t for t, _ in PROJECT_COLUMNS])
    _style_header(ws, len(PROJECT_COLUMNS), 32)

    for p in projects:
        ws.append([date, p["tier"], p["title"], p["client"], p["url"],
                   p["budget"], p["posted"], p["skills"], p["summary"],
                   p["why_fits"], p["proposal_angle"], p["score"], "لم يُقدَّم بعد"])

    widths = [w for _, w in PROJECT_COLUMNS]
    _style_body(ws, len(projects), len(PROJECT_COLUMNS), 60, widths)
    for i in range(len(projects)):
        r = 2 + i
        _link(ws.cell(row=r, column=5))
        score = ws.cell(row=r, column=12)
        score.font = Font(name="Arial", size=11, bold=True, color=ACCENT)
        score.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=2).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True)

    _apply_widths(ws, PROJECT_COLUMNS)
    _dropdown(ws, "M", len(projects), PROJECT_STATUS)
    ws.freeze_panes = "C2"
    if projects:
        ws.auto_filter.ref = f"A1:M{1 + len(projects)}"
    return ws


def _method_sheet(wb, meta, leads):
    ws = wb.create_sheet("المنهجية والمصادر")
    ws.sheet_view.rightToLeft = True

    report = meta.get("report", {})
    rows = [
        ("المنهجية والمصادر", ""),
        ("تاريخ التجهيز", meta.get("date_ar", meta.get("date", ""))),
        ("الأسواق المستهدفة", meta.get("markets", "الخليج")),
        ("مجال اليوم", meta.get("niche", "")),
        ("لغة الرسائل", meta.get("dialect", "لهجة خليجية خفيفة")),
        ("عدد الفرص", str(len(leads))),
        ("", ""),
        ("كيف اُستُخرجت الملاحظات", meta.get("method", "")),
        ("حدود التحقق", meta.get("limits", "")),
        ("حالة المشاريع", meta.get("projects_note", "")),
        ("", ""),
        ("🔥 أفضل ٣ فرص اليوم", "\n".join(report.get("top3", []))),
        ("💰 أفضل فرصة تجارية", report.get("best_commercial", "")),
        ("🎨 أفضل فرصة تصميم", report.get("best_design", "")),
        ("🎥 أفضل فرصة فيديو / UGC", report.get("best_video", "")),
        ("💼 أفضل مشروع Upwork", report.get("best_upwork", "")),
        ("💼 أفضل مشروع مستقل", report.get("best_mostaql", "")),
        ("", ""),
        ("الخطوة التالية المقترحة", meta.get("next_step", "")),
    ]
    for row in rows:
        ws.append(list(row))

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 95
    for r in range(1, len(rows) + 1):
        a, b = ws.cell(row=r, column=1), ws.cell(row=r, column=2)
        a.font = Font(name="Arial", size=10, bold=True, color=INK)
        a.alignment = Alignment(horizontal="right", vertical="top")
        b.font = Font(name="Arial", size=10, color=INK)
        b.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = _auto_height([b.value], [95], 24)
    ws["A1"].font = Font(name="Arial", size=13, bold=True, color=ACCENT)
    ws.row_dimensions[1].height = 24
    return ws


def build(data, out_path):
    meta = data.get("meta", {})
    leads = data["leads"]
    date = meta.get("date", "")

    wb = Workbook()
    _leads_sheet(wb, leads, date)
    _projects_sheet(wb, "Upwork", data.get("upwork", []), date)
    _projects_sheet(wb, "Mostaql", data.get("mostaql", []), date)
    _method_sheet(wb, meta, leads)
    wb.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description="بناء ملف Excel لدفعة قناص")
    ap.add_argument("leads", help="مسار ملف leads.json")
    ap.add_argument("--out", required=True, help="مسار ملف xlsx الناتج")
    args = ap.parse_args()
    print(build(lib.load(args.leads), args.out))


if __name__ == "__main__":
    main()
